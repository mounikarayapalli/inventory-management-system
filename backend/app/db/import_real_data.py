"""Script to import real Calibo AI Academy Inventory data from Excel into PostgreSQL database."""

import openpyxl
from datetime import datetime, date
from decimal import Decimal
from sqlalchemy import text
from app.db.session import SessionLocal
from app.models.role import Role
from app.models.user import User
from app.models.category import Category
from app.models.location import Location
from app.models.supplier import Supplier
from app.models.item import Item

from app.services.user_service import user_service
from app.services.transaction_service import transaction_service
from app.schemas.transaction import (
    OpeningStockRequest,
    InwardRequest,
    OutwardRequest,
)

def import_calibo_data():
    db = SessionLocal()
    print("=== Starting Real Calibo Data Import ===")

    # 1. Clean temporary test data
    print("1. Cleaning temporary test data...")
    db.execute(text("TRUNCATE TABLE stock_movements, inward_transactions, outward_transactions, distribution_transactions, return_transactions, stock_adjustments, opening_stock RESTART IDENTITY CASCADE;"))
    db.execute(text("DELETE FROM items WHERE true;"))
    db.execute(text("DELETE FROM suppliers WHERE true;"))
    db.execute(text("DELETE FROM categories WHERE true;"))
    db.execute(text("DELETE FROM locations WHERE true;"))
    db.execute(text("DELETE FROM users WHERE true;"))
    db.execute(text("DELETE FROM roles WHERE true;"))
    db.commit()

    # 2. Seed Roles & Users
    print("2. Seeding Roles and System Users...")
    admin_role = Role(role_name="admin")
    manager_role = Role(role_name="stock manager")
    db.add_all([admin_role, manager_role])
    db.commit()
    db.refresh(admin_role)
    db.refresh(manager_role)

    admin_user = User(
        username="admin",
        email="admin@calibo.com",
        password_hash=user_service.hash_password("admin123"),
        role_id=admin_role.role_id,
        is_active=True,
    )
    manager_user = User(
        username="manager",
        email="manager@calibo.com",
        password_hash=user_service.hash_password("manager123"),
        role_id=manager_role.role_id,
        is_active=True,
    )
    db.add_all([admin_user, manager_user])
    db.commit()
    db.refresh(admin_user)
    db.refresh(manager_user)

    admin_id = admin_user.user_id

    # 3. Seed Categories (7 Categories)
    print("3. Seeding Categories...")
    category_names = [
        "Bags",
        "Branding",
        "Diaries",
        "Gifts",
        "Goodies",
        "Stationary",
        "T Shirts"
    ]
    categories_map = {}
    for c_name in category_names:
        cat = Category(category_name=c_name)
        db.add(cat)
        db.commit()
        db.refresh(cat)
        categories_map[c_name] = cat.category_id

    # 4. Seed Locations (2 Locations)
    print("4. Seeding Locations...")
    loc_vj = Location(location_name="Vijayawada Hub", description="Vijayawada Academy Warehouse")
    loc_vz = Location(location_name="Vizag Hub", description="Vizag Academy Warehouse")
    db.add_all([loc_vj, loc_vz])
    db.commit()
    db.refresh(loc_vj)
    db.refresh(loc_vz)

    locations_map = {
        "Vijayawada": loc_vj.location_id,
        "Vijayawada Hub": loc_vj.location_id,
        "Vizag": loc_vz.location_id,
        "Vizag Hub": loc_vz.location_id,
    }

    # 5. Seed Suppliers (5 Suppliers)
    print("5. Seeding Suppliers...")
    supplier_names = [
        "Red Chariot",
        "Local Print",
        "Local Print (Vizag)",
        "Local Print (Vijayawada)",
        "Subin P"
    ]
    suppliers_map = {}
    for s_name in supplier_names:
        sup = Supplier(supplier_name=s_name, email=f"{s_name.lower().replace(' ', '').replace('(', '').replace(')', '')}@vendor.com")
        db.add(sup)
        db.commit()
        db.refresh(sup)
        suppliers_map[s_name] = sup.supplier_id

    suppliers_map["Local Print, Vizag"] = suppliers_map["Local Print (Vizag)"]
    suppliers_map["Local Print, Vijayawada"] = suppliers_map["Local Print (Vijayawada)"]

    # 6. Seed Items (17 SKUs)
    print("6. Seeding Item Catalog (17 SKUs)...")
    items_catalog = [
        ("S101", "Calibo Badges", "Branding", "pcs"),
        ("S102", "Calibo Branded Folders", "Branding", "pcs"),
        ("S103", "Calibo Hangings - 3x8", "Branding", "pcs"),
        ("S104", "Calibo Hangings+ KL", "Branding", "pcs"),
        ("S105", "College Brouchures", "Stationary", "pcs"),
        ("S106", "Diaries+Pens", "Diaries", "pcs"),
        ("S107", "Key Chains with Calibo Logo", "Goodies", "pcs"),
        ("S108", "Laptop Bags + Water Bottles", "Bags", "set"),
        ("S109", "Lepakshi Gifts", "Gifts", "pcs"),
        ("S110", "Staff T Shirts", "T Shirts", "pcs"),
        ("S111", "Standees Large", "Branding", "pcs"),
        ("S112", "Standees KL Branding", "Branding", "pcs"),
        ("S113", "Student Brouchures", "Stationary", "pcs"),
        ("S114", "Student T Shirts", "T Shirts", "pcs"),
        ("S115", "Table stand Flags", "Branding", "pcs"),
        ("S116", "USB Cables", "Goodies", "pcs"),
        ("S117", "Vendor T Shirts", "T Shirts", "pcs"),
    ]

    items_map = {}
    for code, name, cat_name, unit in items_catalog:
        cat_id = categories_map[cat_name]
        it = Item(
            item_code=code,
            item_name=name,
            category_id=cat_id,
            unit=unit,
            minimum_level=5,
            default_unit_cost=Decimal("0.00"),
        )
        db.add(it)
        db.commit()
        db.refresh(it)
        items_map[code] = it.item_id
        items_map[name] = it.item_id

    print(f"Master Data Seed Complete! Categories: {len(categories_map)}, Locations: 2, Suppliers: {len(supplier_names)}, Items: {len(items_catalog)}")

    # 7. Parse and Import Excel Transactions
    excel_path = r'C:\Users\ADMIN\.gemini\antigravity\brain\5bafea04-1cf2-4756-8b77-9b317ac063d0\scratch\temp_excel.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    opening_count = 0
    inward_count = 0
    outward_count = 0
    unimported_rows = []
    recorded_openings = set()

    print("\n7. Importing Transactions from Excel...")

    for row_idx in range(2, sheet.max_row + 1):
        sl = sheet.cell(row=row_idx, column=1).value
        hub_str = sheet.cell(row=row_idx, column=2).value
        item_name = sheet.cell(row=row_idx, column=3).value
        code = sheet.cell(row=row_idx, column=4).value
        cat = sheet.cell(row=row_idx, column=5).value
        op_qty = sheet.cell(row=row_idx, column=6).value
        dt_rec = sheet.cell(row=row_idx, column=7).value
        rec_from = sheet.cell(row=row_idx, column=8).value
        qty_rec = sheet.cell(row=row_idx, column=9).value
        dt_ut = sheet.cell(row=row_idx, column=10).value
        iss_to = sheet.cell(row=row_idx, column=11).value
        purpose = sheet.cell(row=row_idx, column=12).value
        qty_ut = sheet.cell(row=row_idx, column=13).value
        bal = sheet.cell(row=row_idx, column=14).value
        remarks = sheet.cell(row=row_idx, column=16).value

        if not hub_str or not (code or item_name):
            continue

        loc_id = locations_map.get(hub_str, loc_vj.location_id)
        item_id = items_map.get(code) or items_map.get(item_name)

        if not item_id:
            unimported_rows.append((row_idx, f"Unknown item code {code} / name {item_name}"))
            continue

        # Parse unit cost from remarks if present
        unit_cost = Decimal("0.00")
        if remarks and "Cost Rs." in str(remarks):
            try:
                rem_str = str(remarks)
                cost_part = rem_str.split("Cost Rs.")[1].strip().split()[0].replace(',', '')
                tot_cost = Decimal(cost_part)
                parsed_qty = Decimal(str(qty_rec).replace('+', '')) if qty_rec else Decimal("1")
                if parsed_qty > 0:
                    unit_cost = tot_cost / parsed_qty
            except Exception:
                unit_cost = Decimal("0.00")

        # A. Opening Stock
        if op_qty is not None and str(op_qty).strip() != "":
            try:
                clean_op = str(op_qty).replace('+', '').strip()
                parsed_op = Decimal(clean_op)
                op_key = (item_id, loc_id)
                if parsed_op > 0 and op_key not in recorded_openings:
                    op_req = OpeningStockRequest(
                        item_id=item_id,
                        location_id=loc_id,
                        quantity=parsed_op,
                        unit_cost=unit_cost,
                        opening_date=date(2026, 4, 1),
                        remarks=f"Initial Opening Stock from Register Row {row_idx}"
                    )
                    transaction_service.record_opening_stock(db, payload=op_req, created_by=admin_id)
                    recorded_openings.add(op_key)
                    opening_count += 1
            except Exception as e:
                pass

        # B. Inward Receipt
        if qty_rec is not None and str(qty_rec).strip() != "":
            try:
                clean_qty_str = str(qty_rec).replace('+', '').strip()
                parsed_qty_rec = Decimal(clean_qty_str)

                if parsed_qty_rec > 0:
                    sup_id = suppliers_map.get(rec_from) if rec_from else suppliers_map.get("Subin P")
                    if not sup_id:
                        sup_id = suppliers_map.get("Subin P")

                    rec_date = date.today()
                    if isinstance(dt_rec, datetime):
                        rec_date = dt_rec.date()
                    elif isinstance(dt_rec, date):
                        rec_date = dt_rec

                    inward_req = InwardRequest(
                        inward_no=f"INW-EXCEL-R{row_idx:02d}",
                        item_id=item_id,
                        location_id=loc_id,
                        supplier_id=sup_id,
                        quantity=parsed_qty_rec,
                        unit_cost=unit_cost,
                        invoice_no=f"EXCEL-R{row_idx:02d}",
                        inward_date=rec_date,
                        remarks=f"{purpose or ''} - {remarks or ''}".strip(" -")
                    )
                    transaction_service.record_inward(db, payload=inward_req, created_by=admin_id)
                    inward_count += 1
            except Exception as e:
                unimported_rows.append((row_idx, f"Inward error: {e}"))

        # C. Outward Issue
        if qty_ut is not None and str(qty_ut).strip() != "" and str(qty_ut).strip().lower() != "reusable":
            try:
                clean_ut_str = str(qty_ut).strip()
                parsed_qty_ut = Decimal(clean_ut_str)

                if parsed_qty_ut > 0:
                    ut_date = date.today()
                    if isinstance(dt_ut, datetime):
                        ut_date = dt_ut.date()
                    elif isinstance(dt_ut, date):
                        ut_date = dt_ut

                    outward_req = OutwardRequest(
                        outward_no=f"OUT-EXCEL-U{row_idx:02d}",
                        item_id=item_id,
                        location_id=loc_id,
                        quantity=parsed_qty_ut,
                        reference_no=f"EXCEL-U{row_idx:02d}",
                        issued_to=str(iss_to) if iss_to else "Event Utilisation",
                        purpose=str(purpose) if purpose else "Calibo Academy Event",
                        outward_date=ut_date,
                        remarks=f"{remarks or ''}".strip()
                    )
                    transaction_service.record_outward(db, payload=outward_req, created_by=admin_id)
                    outward_count += 1
            except Exception as e:
                unimported_rows.append((row_idx, f"Outward error: {e}"))

    db.close()

    print(f"\n=== Import Completed Successfully ===")
    print(f"Opening Stock Entries: {opening_count}")
    print(f"Inward Transactions Imported: {inward_count}")
    print(f"Outward Transactions Imported: {outward_count}")
    print(f"Unimported / Skipped Rows: {len(unimported_rows)}")
    if unimported_rows:
        for r_idx, err in unimported_rows:
            print(f"  Row {r_idx:02d}: {err}")

if __name__ == "__main__":
    import_calibo_data()
